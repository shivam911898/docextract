"""DocExtract core: extraction + validation logic shared by the Streamlit app and the CLI.

Pipeline per file:  ingest (PDF/PNG/JPG, or a .zip of them)  ->  extract (Groq text model for text, Gemini vision for
scans and photos, or regex/heuristic without a key; scans and photos reach the text routes through OCR: Docling if
installed, else the lightweight RapidOCR / Tesseract in ocr.py)  ->  validate (arithmetic cross-checks)  ->  flat row
+ line items for CSV / XLSX / Google Sheets.
"""
from __future__ import annotations

import html
import io
import json
import os
import re
import zipfile
from datetime import date as _date
from pathlib import Path, PurePosixPath
from typing import List, Optional, Tuple

import pandas as pd
from pydantic import BaseModel, Field

try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:  # pragma: no cover
    pass

import llm  # noqa: E402  (after load_dotenv so .env is already in the environment)
import ocr  # noqa: E402

# ---------------------------------------------------------------- settings
TOLERANCE = 0.005          # relative tolerance for arithmetic checks (0.5%) ...
ABS_TOLERANCE = 0.05       # ... or this many currency units, whichever is larger (rounding noise)
MIN_CONFIDENCE = 0.7
OCR_CONFIDENCE_CAP = 0.85  # text that went through OCR can have misread digits: confidence never exceeds this
IMG_EXTS = {".png", ".jpg", ".jpeg"}
EXTS = {".pdf"} | IMG_EXTS
DOC_TYPES = ("invoice", "receipt", "other")
MAX_FILES = int(os.getenv("MAX_FILES", "10"))              # per UI batch (protects a public demo's API quota); CLI is uncapped
MAX_PAGES = int(os.getenv("MAX_PAGES", "3"))               # pages sent to the vision model per document
MAX_FILE_BYTES = int(os.getenv("MAX_FILE_MB", "15")) * 1024 * 1024
MAX_ZIP_BYTES = 150 * 1024 * 1024                          # total uncompressed size of one .zip
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite")
GEMINI_RETRY_ATTEMPTS = 5
GEMINI_RETRY_DELAY = 4     # seconds; the SDK backs off exponentially from here
CURRENCIES = {"USD", "EUR", "GBP", "CAD", "AUD", "NZD", "CHF", "SEK", "NOK", "DKK", "INR", "JPY"}


class LineItem(BaseModel):
    description: str = ""
    qty: Optional[float] = None
    unit_price: Optional[float] = None
    amount: Optional[float] = None


class Document(BaseModel):
    """One invoice or receipt. Also the response schema handed to Gemini."""
    doc_type: Optional[str] = Field(None, description="One of: invoice, receipt, other")
    vendor: Optional[str] = Field(None, description="Seller / merchant name")
    document_number: Optional[str] = Field(None, description="Invoice number, or receipt / check / transaction number")
    date: Optional[str] = Field(None, description="Issue / purchase date, ISO YYYY-MM-DD if possible")
    currency: Optional[str] = Field(None, description="3-letter ISO code")
    line_items: List[LineItem] = Field(default_factory=list)
    subtotal: Optional[float] = None
    tax: Optional[float] = Field(None, description="Total tax / VAT / GST amount")
    tip: Optional[float] = Field(None, description="Tip / gratuity / service charge, if printed")
    total: Optional[float] = Field(None, description="Final amount due / paid")
    payment_method: Optional[str] = Field(None, description="Receipts: how it was paid, e.g. 'Visa ****4242', 'cash'")
    confidence: float = Field(0.5, ge=0, le=1)


# ---------------------------------------------------------------- file helpers
def has_gemini_key() -> bool:
    return bool(os.getenv("GEMINI_API_KEY"))


def active_provider() -> str:
    """'groq', 'gemini' or 'none' (see llm.provider for the rules)."""
    return llm.provider()


def ocr_label() -> str:
    """The OCR that reads scans and photos for the text routes: 'Docling' (if installed, it wins), 'RapidOCR',
    'Tesseract' (see ocr.py and OCR_BACKEND) or 'off'."""
    return "Docling" if docling_available() else ocr.label()


def ocr_available() -> bool:
    return docling_available() or ocr.available()


def provider_label() -> str:
    """One line for the UI / CLI: which provider and models are in use, and which OCR reads scans and photos."""
    p = active_provider()
    name = ocr_label()
    scans = (f"scans and photos read by OCR ({name})" if name != "off"
             else "scans and photos need Gemini or an OCR backend")
    if p == "groq":
        return f"Groq: {llm.groq_text_model()} (text-layer PDFs; {scans})"
    if p == "gemini":
        return f"Gemini: {GEMINI_MODEL} (reads PDFs, scans and photos with vision)"
    return f"none (offline regex / heuristic extractor; {scans})"


def _png(img) -> bytes:
    buf = io.BytesIO(); img.convert("RGB").save(buf, "PNG")
    return buf.getvalue()


def render_pages(data: bytes, filename: str, max_pages: int = MAX_PAGES) -> List[bytes]:
    """PNG bytes for the first `max_pages` pages of a PDF, or the image itself (one page)."""
    if Path(filename).suffix.lower() == ".pdf":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(data)
        return [_png(pdf[i].render(scale=2).to_pil()) for i in range(min(len(pdf), max_pages))]
    from PIL import Image, ImageOps
    return [_png(ImageOps.exif_transpose(Image.open(io.BytesIO(data))))]      # phone photos: honour the EXIF rotation


def render_first_page(data: bytes, filename: str) -> bytes:
    return render_pages(data, filename, 1)[0]


def pdf_text(data: bytes, filename: str) -> str:
    """Text layer of a PDF in layout mode (keeps table columns on one line). '' for images / scans."""
    if Path(filename).suffix.lower() != ".pdf":
        return ""
    from pypdf import PdfReader
    try:
        pages = PdfReader(io.BytesIO(data)).pages[:MAX_PAGES]
        out = []
        for p in pages:
            try:
                out.append(p.extract_text(extraction_mode="layout") or "")
            except Exception:
                out.append(p.extract_text() or "")
        return "\n".join(out)
    except Exception:
        return ""


_docling_converter = None
_docling_warned = False


def _md_to_lines(md: str) -> str:
    """Flatten Docling markdown (headings, tables) into the one-row-per-line text the heuristic parser expects."""
    out: List[str] = []
    pending: List[int] = []                            # label lines still waiting for their amount
    for l in (x.strip() for x in md.splitlines()):
        if not l or l.startswith("<!--") or re.fullmatch(r"[\s:|-]+", l):
            continue
        l = html.unescape(re.sub(r"^#+\s*", "", l))
        if l.startswith("|"):
            l = "   ".join(c.strip() for c in l.strip("|").split("|") if c.strip())
        if re.fullmatch(rf"{SYM}{NUM}", l):            # lone amount: Docling often emits "Subtotal / VAT / amt / amt"
            if pending:
                out[pending.pop(0)] += "   " + l
                continue
            if out and not re.search(NUM + r"\s*$", out[-1]):
                out[-1] += "   " + l
                continue
        out.append(l)
        if LABELS.match(l) and not re.search(NUM, l) and not re.match(r"(?i)tax\s+invoice", l):
            pending.append(len(out) - 1)
    return "\n".join(out)


def docling_available() -> bool:
    """True when the optional Docling OCR backend (requirements-docling.txt) is installed."""
    import importlib.util
    return importlib.util.find_spec("docling") is not None


def docling_text(data: bytes, filename: str) -> str:
    """Optional backend (EXTRACT_BACKEND=docling): layout + table model, and OCR for images. Slow on CPU."""
    global _docling_converter
    from docling.datamodel.base_models import DocumentStream
    from docling.document_converter import DocumentConverter
    _docling_converter = _docling_converter or DocumentConverter()
    res = _docling_converter.convert(DocumentStream(name=Path(filename).name, stream=io.BytesIO(data)))
    return _md_to_lines(res.document.export_to_markdown())


def ocr_text(data: bytes, filename: str) -> Tuple[str, Optional[float]]:
    """OCR with the lightweight backend (ocr.py): PDF pages are rendered with pypdfium2, images are decoded as they
    are. Returns (text, mean OCR confidence 0-1)."""
    pages = render_pages(data, filename) if Path(filename).suffix.lower() == ".pdf" else [data]
    return ocr.read_pages(pages)


def scan_text(data: bytes, filename: str) -> Tuple[str, Optional[float]]:
    """Text of a file that has no text layer (photo, image, scanned PDF): Docling if installed, else the lightweight
    OCR, else ScannedImageError (the Needs review row). Returns (text, OCR confidence or None)."""
    if docling_available():
        text, score = docling_text(data, filename), None
    elif ocr.available():
        text, score = ocr_text(data, filename)
    else:
        raise ScannedImageError(SCANNED_MSG)
    if not _usable(text):
        raise ValueError("no readable text found, even with OCR")
    return text, score


def ocr_confidence(conf: float, ocr_score: Optional[float] = None) -> float:
    """Confidence of a document read from OCR text: capped at OCR_CONFIDENCE_CAP and never above how sure the OCR
    engine itself was (a blurry photo that OCR struggled with lands under MIN_CONFIDENCE, so it is flagged)."""
    return round(min(conf, OCR_CONFIDENCE_CAP, 1.0 if ocr_score is None else ocr_score), 2)


def doc_text(data: bytes, filename: str) -> str:
    """Text for the heuristic parser / model prompt, from the configured backend (falls back to pypdf)."""
    global _docling_warned
    if os.getenv("EXTRACT_BACKEND", "pypdf").strip().lower() == "docling":
        try:
            return docling_text(data, filename)
        except Exception as e:
            if not _docling_warned:
                _docling_warned = True
                print(f"WARNING: EXTRACT_BACKEND=docling failed ({type(e).__name__}: {str(e)[:120]}); using pypdf.")
    return pdf_text(data, filename)


def expand_uploads(items: List[Tuple[str, bytes]], max_files: Optional[int] = None) -> Tuple[List[Tuple[str, bytes]], List[str]]:
    """Flatten uploads: pass PDFs/images through, unpack .zip files (in memory, never written to disk).

    Returns (documents, notes). Notes describe anything skipped (wrong type, too big, batch cap).
    `max_files` caps the batch (the UI passes MAX_FILES); None = uncapped (the CLI).
    """
    docs: List[Tuple[str, bytes]] = []
    notes: List[str] = []

    def add(name: str, data: bytes):
        if Path(name).suffix.lower() not in EXTS:
            notes.append(f"{name}: skipped (unsupported type)")
        elif len(data) > MAX_FILE_BYTES:
            notes.append(f"{name}: skipped (over {MAX_FILE_BYTES // 1024 // 1024} MB)")
        elif not data:
            notes.append(f"{name}: skipped (empty file)")
        else:
            docs.append((name, data))

    for name, data in items:
        if Path(name).suffix.lower() != ".zip":
            add(name, data)
            continue
        try:
            zf = zipfile.ZipFile(io.BytesIO(data))
            infos = [i for i in zf.infolist() if not i.is_dir()]
            if sum(i.file_size for i in infos) > MAX_ZIP_BYTES:
                notes.append(f"{name}: skipped (zip expands to more than {MAX_ZIP_BYTES // 1024 // 1024} MB)")
                continue
            for i in infos:
                parts = [p for p in PurePosixPath(i.filename.replace("\\", "/")).parts if p not in ("/", "..", ".")]
                if not parts or parts[-1].startswith(".") or "__MACOSX" in parts:
                    continue
                add("/".join(parts), zf.read(i))
        except zipfile.BadZipFile:
            notes.append(f"{name}: skipped (not a valid zip)")
    if max_files is not None and len(docs) > max_files:
        notes.append(f"Only the first {max_files} of {len(docs)} files were processed: this demo caps a batch at "
                     f"{max_files} files to stay inside the free API quota. Upload the rest in a second batch, or "
                     "run the app yourself (MAX_FILES) or the CLI for bigger jobs.")
        docs = docs[:max_files]
    return docs, notes


# ---------------------------------------------------------------- Gemini
PROMPT = (
    "Extract this business document (an invoice or a receipt) into the JSON schema. "
    "doc_type: 'invoice', 'receipt', or 'other' if it is neither. "
    "vendor is the seller / merchant; document_number is the invoice number or the receipt / check / transaction "
    "number; payment_method is how a receipt was paid (e.g. 'Visa ****4242', 'cash'); tip only if printed. "
    "Use null for anything not visible. Dates as YYYY-MM-DD (read 03/04/2025 as MM/DD for US documents and DD/MM "
    "for UK/EU ones), currency as 3-letter ISO code, numbers as plain floats (no symbols). "
    "Include every line item. Set confidence (0-1) to how sure you are that all values are read correctly. "
    "Extracted PDF text (may be empty or noisy) for reference:\n"
)


def extract_gemini(pngs: List[bytes], text: str) -> Document:
    from google import genai
    from google.genai import types
    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"], http_options=genai.types.HttpOptions(base_url=os.getenv("GEMINI_BASE_URL") or None, timeout=90000, retry_options=genai.types.HttpRetryOptions(attempts=GEMINI_RETRY_ATTEMPTS, initial_delay=GEMINI_RETRY_DELAY, http_status_codes=[429, 500, 502, 503, 504])))
    parts = [types.Part.from_bytes(data=p, mime_type="image/png") for p in pngs]
    resp = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=parts + [PROMPT + text[:6000]],
        config=types.GenerateContentConfig(response_mime_type="application/json", response_schema=Document, temperature=0),
    )
    return Document.model_validate_json(resp.text)


# ---------------------------------------------------------------- Groq (OpenAI-compatible, JSON mode)
MIN_TEXT_CHARS = 40        # fewer non-space characters than this = no usable text layer (scan / photo)
GROQ_TEXT_CHARS = 8000     # text sent to the model (keeps a request inside the free tier's tokens-per-minute)
SCHEMA_HINT = (
    '{"doc_type": "invoice" | "receipt" | "other", "vendor": string | null, "document_number": string | null, '
    '"date": "YYYY-MM-DD" | null, "currency": 3-letter ISO code string such as "USD", "EUR", "GBP", "SEK" | null, "line_items": [{"description": string, "qty": number | '
    'null, "unit_price": number | null, "amount": number | null}], "subtotal": number | null, "tax": number | null, '
    '"tip": number | null, "total": number | null, "payment_method": string | null, "confidence": number}'
)
GROQ_PROMPT = (
    "Extract this business document (an invoice or a receipt) into JSON. Reply with ONE JSON object and nothing "
    "else, using exactly these keys:\n" + SCHEMA_HINT + "\n"
    "doc_type: 'invoice', 'receipt', or 'other' if it is neither. "
    "vendor is the seller / merchant; document_number is the invoice number or the receipt / check / transaction "
    "number; payment_method is how a receipt was paid (e.g. 'Visa ****4242', 'cash'); tip only if printed. "
    "Use null for anything not visible. Dates as YYYY-MM-DD (read 03/04/2025 as MM/DD for US documents and DD/MM "
    "for UK/EU ones), currency as 3-letter ISO code, numbers as plain numbers (no currency symbols, no thousands "
    "separators). Include every line item. Never invent or recompute values: copy what is printed. confidence is "
    "a number from 0 to 1 for how sure you are that every value is read correctly.\n"
)


OCR_NOTE = (
    "The text below was read from an image by OCR, so it can contain mistakes: words may be run together (for "
    "example 'MAINSTREET CAFE' for 'MAIN STREET CAFE') and look-alike characters get confused (O/0, l/I/1). Restore "
    "the missing spaces in names and descriptions, but copy every number exactly as it appears and never guess a "
    "digit.\n"
)


def _doc_from_json(content: str) -> Document:
    """Parse a model reply into a Document. Tolerates code fences and chatter around the object and the usual
    model quirks (null line_items, 0-100 confidence); anything else raises, which becomes a Needs review row."""
    m = re.search(r"\{.*\}", content, re.S)
    if not m:
        raise ValueError(f"model reply contained no JSON object: {content[:80]!r}")
    try:
        obj = json.loads(m.group(0))
    except json.JSONDecodeError as e:
        raise ValueError(f"model returned malformed JSON ({e.msg} at char {e.pos})") from None
    if not isinstance(obj, dict):
        raise ValueError("model JSON is not an object")
    obj["line_items"] = [{**li, "description": li.get("description") or ""}
                         for li in (obj.get("line_items") or []) if isinstance(li, dict)]
    c = obj.get("confidence")
    if c is None:
        obj.pop("confidence", None)
    elif isinstance(c, (int, float)) and 1 < c <= 100:
        obj["confidence"] = c / 100
    return Document.model_validate(obj)


SCANNED_MSG = "scanned image: needs a vision provider (Gemini) or an OCR backend (RapidOCR, Tesseract or Docling)"


class ScannedImageError(ValueError):
    """The file has no text layer and no OCR / vision route is available for the active provider."""


def _usable(text: str) -> bool:
    return len("".join(text.split())) >= MIN_TEXT_CHARS


def extract_groq(data: bytes, filename: str, text: str) -> Tuple[Document, str, str, Optional[float]]:
    """Groq reads text only. A text-layer document goes to the text model as is. An image or scanned PDF (no usable
    text layer) is OCR'd first (scan_text: Docling if installed, else RapidOCR / Tesseract); with no OCR backend it
    cannot be read here and becomes a Needs review row. Returns (document, model id, source, OCR confidence) where
    source is 'text' (text layer) or 'ocr'."""
    source, score = "text", None
    if not _usable(text):
        text, score = scan_text(data, filename)
        source = "ocr"
    model = llm.groq_text_model()
    prompt = GROQ_PROMPT + (OCR_NOTE if source == "ocr" else "") + "\nDocument text:\n" + text[:GROQ_TEXT_CHARS]
    reply = llm.groq_chat_json([{"role": "user", "content": prompt}], model)
    return _doc_from_json(reply), model, source, score


# ---------------------------------------------------------------- heuristic (no key)
NUM = r"-?\d[\d,]*\.\d{2}"
SYM = r"[^\w\s.,-]{0,3}\s?"        # optional currency symbol glued to a number
SYMBOLS = {"$": "USD", "€": "EUR", "£": "GBP", "₹": "INR"}
MONTHS = {m: i for i, m in enumerate("jan feb mar apr may jun jul aug sep oct nov dec".split(), 1)}
MON = r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?"
DATE_RE = re.compile("|".join([r"\d{4}-\d{2}-\d{2}", r"\d{1,2}[/.]\d{1,2}[/.]\d{2,4}",
                               r"\d{1,2}\s+" + MON + r"\s+\d{4}", MON + r"\s+\d{1,2},?\s+\d{4}"]))
L_SUB = r"sub\s*-?\s*total|net\s+total|net\s+amount"
L_TAX = r"(?:(?:incl\w*|includes?|of\s+which)\s+)?(?:sales\s+tax|tax|vat|gst|mwst)"
L_TIP = r"tip|gratuity|service\s+charge"
L_TOTAL = r"grand\s+total|total\s+due|amount\s+due|balance\s+due|total\s+amount|total"
LABELS = re.compile(rf"(?i)^\s*(?:{L_SUB}|{L_TAX}|{L_TIP}|{L_TOTAL})\b")
NOT_ITEM = re.compile(r"(?i)^(?:sub\s*-?\s*total|total|tax|vat|gst|tip|description|item|date|invoice|receipt|payment|paid|card|tel|balance|due|bill)\b")


def _num(s: str) -> float:
    return float(s.replace(",", ""))


def _norm(line: str) -> str:
    return re.sub(r"\s+", " ", line).strip()


def _label_amount(lines: List[str], label: str) -> Optional[float]:
    """Last amount on the last line that starts with `label` (e.g. 'Sales Tax 8.0% 1.71' -> 1.71)."""
    hit = None
    for l in lines:
        if re.match(rf"(?i)^(?:{label})\b", l):
            nums = re.findall(NUM, l)
            if nums:
                hit = _num(nums[-1])
    return hit


def _to_iso(s: str, currency: Optional[str] = None) -> Optional[str]:
    """Parse the date formats seen on invoices/receipts into YYYY-MM-DD; None if unparseable."""
    s = s.strip()
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
            return _date.fromisoformat(s).isoformat()
        m = re.fullmatch(r"(\d{1,2})([/.])(\d{1,2})\2(\d{2,4})", s)
        if m:
            a, sep, b, y = int(m[1]), m[2], int(m[3]), int(m[4])
            y += 2000 if y < 100 else 0
            if sep == "." or a > 12:
                d, mo = a, b                      # day first
            elif b > 12:
                mo, d = a, b                      # month first (US)
            else:                                 # ambiguous: US/CA are month-first, everyone else day-first
                mo, d = (a, b) if currency in ("USD", "CAD") else (b, a)
            return _date(y, mo, d).isoformat()
        m = re.fullmatch(r"(\d{1,2})\s+([A-Za-z]+)\.?\s+(\d{4})", s)
        if m:
            return _date(int(m[3]), MONTHS[m[2][:3].lower()], int(m[1])).isoformat()
        m = re.fullmatch(r"([A-Za-z]+)\.?\s+(\d{1,2}),?\s+(\d{4})", s)
        if m:
            return _date(int(m[3]), MONTHS[m[1][:3].lower()], int(m[2])).isoformat()
    except (ValueError, KeyError):
        pass
    return None


def _find_date(lines: List[str], currency: Optional[str]) -> Optional[str]:
    for l in lines:                                    # prefer a labelled date, never a due date
        if re.search(r"(?i)\bdate\b", l) and not re.search(r"(?i)\bdue\b", l):
            m = DATE_RE.search(l)
            if m and (iso := _to_iso(m.group(0), currency)):
                return iso
    for l in lines:
        if not re.search(r"(?i)\bdue\b", l):
            m = DATE_RE.search(l)
            if m and (iso := _to_iso(m.group(0), currency)):
                return iso
    return None


def _find_currency(text: str) -> Optional[str]:
    m = re.search(r"(?i)\b(?:currency|amounts?\s+in)\s*:?\s*([A-Za-z]{3})\b", text)
    if m and m.group(1).upper() in CURRENCIES:
        return m.group(1).upper()
    codes = {c: len(re.findall(rf"\b{c}\b", text)) for c in CURRENCIES}
    codes = {c: n for c, n in codes.items() if n}
    if codes:
        return max(codes, key=codes.get)
    syms = {c: text.count(s) for s, c in SYMBOLS.items() if s in text}
    return max(syms, key=syms.get) if syms else None


def _detect_type(text: str) -> str:
    inv, rec = re.search(r"(?i)\binvoice\b", text), re.search(r"(?i)\breceipt\b", text)
    if inv and rec:
        return "invoice" if inv.start() < rec.start() else "receipt"
    if inv:
        return "invoice"
    if rec or re.search(r"(?i)\b(visa|mastercard|amex|contactless|change due|thank you for shopping)\b", text):
        return "receipt"
    return "other"


def _parse_items(lines: List[str]) -> List[LineItem]:
    """Line items = lines above the first subtotal / tax / total row."""
    stop = next((i for i, l in enumerate(lines) if LABELS.match(l) and re.search(NUM, l)), len(lines))
    items: List[LineItem] = []
    for l in lines[:stop]:
        if NOT_ITEM.match(l):
            continue
        if m := re.match(rf"^(?P<d>.*?[A-Za-z].*?)\s+(?P<q>\d+(?:\.\d+)?)\s+{SYM}(?P<u>{NUM})\s+{SYM}(?P<a>{NUM})$", l):
            items.append(LineItem(description=m["d"], qty=_num(m["q"]), unit_price=_num(m["u"]), amount=_num(m["a"])))
        elif m := re.match(rf"^(?P<d>.*?[A-Za-z].*?)\s+(?P<q>\d+)\s*@\s*{SYM}(?P<u>{NUM})\s+{SYM}(?P<a>{NUM})$", l):
            items.append(LineItem(description=m["d"], qty=_num(m["q"]), unit_price=_num(m["u"]), amount=_num(m["a"])))
        elif m := re.match(rf"^(?P<q>\d+)\s*x\s+(?P<d>.*?[A-Za-z].*?)\s+{SYM}(?P<a>{NUM})$", l):
            q, a = _num(m["q"]), _num(m["a"])
            items.append(LineItem(description=m["d"], qty=q, unit_price=round(a / q, 2) if q else None, amount=a))
        elif m := re.match(rf"^(?P<q>\d{{1,2}})\s+(?P<d>[A-Za-z].*?)\s+{SYM}(?P<a>{NUM})$", l):   # "2 Iced Tea 7.50"
            q, a = _num(m["q"]), _num(m["a"])
            items.append(LineItem(description=m["d"], qty=q, unit_price=round(a / q, 2) if q else None, amount=a))
        elif m := re.match(rf"^(?P<d>.*?[A-Za-z].*?)\s+{SYM}(?P<a>{NUM})$", l):
            items.append(LineItem(description=m["d"], amount=_num(m["a"])))
    return items


def extract_heuristic(text: str) -> Document:
    """Regex/heuristic extractor over a text-layer PDF. Deliberately conservative: confidence is capped at 0.72."""
    raw = [l for l in text.splitlines() if l.strip()]
    lines = [_norm(l) for l in raw]
    flat = "\n".join(lines)
    doc = Document(confidence=0.0, doc_type=_detect_type(flat))
    for l in raw:                                      # vendor = first cell of the first meaningful line
        cell = re.split(r"\s{3,}", l.strip())[0]
        if re.search(r"[A-Za-z]{2}", cell) and not re.fullmatch(r"(?i)(tax\s+)?(invoice|receipt)", cell):
            doc.vendor = cell.title() if re.fullmatch(r"[A-Z0-9 &.,'-]+", cell) else cell
            break
    m = re.search(r"(?i)\b(?:invoice|receipt|check|ticket)\s*(?:number|no|nr)\b\.?\s*[:#.]?\s*([A-Z0-9][A-Z0-9\-/]*\d[A-Z0-9\-/]*)", flat) \
        or re.search(r"(?i)\b(?:invoice|receipt|check|ticket)\s*#\s*[:.]?\s*([A-Z0-9][A-Z0-9\-/]*)", flat)
    doc.document_number = m.group(1) if m else None
    doc.currency = _find_currency(flat)
    doc.date = _find_date(lines, doc.currency)
    doc.subtotal = _label_amount(lines, L_SUB)
    doc.tax = _label_amount(lines, L_TAX)
    doc.tip = _label_amount(lines, L_TIP)
    doc.total = _label_amount(lines, L_TOTAL)
    doc.line_items = _parse_items(lines)
    if doc.doc_type != "invoice":
        m = re.search(r"(?im)^(?:paid(?:\s+(?:by|with))?|payment(?:\s+method)?|card|tendered)\s*:\s*(.+)$", flat) \
            or re.search(r"(?i)\b(visa|mastercard|amex|american express|apple pay|google pay|paypal|cash)\b.*", flat)
        doc.payment_method = _norm(m.group(1) if m.lastindex else m.group(0)) if m else None
    found = sum(x is not None for x in (doc.vendor, doc.document_number, doc.date, doc.currency, doc.subtotal,
                                          doc.tax, doc.total)) + bool(doc.line_items)
    doc.confidence = round(min(0.72, 0.1 + 0.1 * found), 2)  # capped below vision-model confidence on purpose
    return doc


# ---------------------------------------------------------------- validation
def _close(a: float, b: float) -> bool:
    return abs(a - b) <= max(ABS_TOLERANCE, TOLERANCE * max(abs(a), abs(b)))


def normalise(doc: Document) -> Document:
    """Clean model / heuristic output: valid doc_type, ISO currency, ISO date."""
    t = (doc.doc_type or "").strip().lower()
    doc.doc_type = t if t in DOC_TYPES else ("other" if t else None)
    if doc.currency:
        doc.currency = doc.currency.strip().upper()[:3] or None
    if doc.date and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", doc.date):
        doc.date = _to_iso(doc.date, doc.currency) or doc.date
    return doc


def validate(doc: Document) -> List[str]:
    """Reasons a human should look at this document ([] = looks consistent)."""
    reasons = []
    amts = [li.amount for li in doc.line_items if li.amount is not None]
    tip = doc.tip or 0
    if doc.doc_type not in ("invoice", "receipt"):
        reasons.append("not recognised as an invoice or receipt")
    if doc.subtotal is not None and amts and not _close(sum(amts), doc.subtotal):
        reasons.append(f"line items sum {sum(amts):.2f} != subtotal {doc.subtotal:.2f}")
    if doc.subtotal is not None and doc.total is not None:
        exp = doc.subtotal + (doc.tax or 0) + tip
        if not _close(exp, doc.total):
            reasons.append(f"subtotal+tax{'+tip' if tip else ''} {exp:.2f} != total {doc.total:.2f}")
    elif doc.subtotal is None and doc.total is not None and amts:   # e.g. UK receipts: prices already include VAT
        s = sum(amts)
        if not (_close(s + (doc.tax or 0) + tip, doc.total) or _close(s + tip, doc.total)):
            reasons.append(f"line items sum {s:.2f} (+tax/tip) != total {doc.total:.2f}")
    if doc.total is None:
        reasons.append("total missing")
    if not doc.line_items:
        reasons.append("no line items found")
    if not doc.vendor:
        reasons.append("vendor missing")
    if not doc.date:
        reasons.append("date missing")
    if doc.confidence < MIN_CONFIDENCE:
        reasons.append(f"low confidence {doc.confidence:.2f}")
    return reasons


def heuristic_from_file(name: str, data: bytes) -> Tuple[Document, str, Optional[float]]:
    """No-key extraction. With EXTRACT_BACKEND=docling, text-layer PDFs are parsed from both the Docling text and the
    pypdf text and the parse that validates cleaner wins (Docling adds table structure). A file with no text layer
    (photo, scan) is OCR'd first (scan_text); without an OCR backend it raises ScannedImageError.
    Returns (document, source 'text' | 'ocr', OCR confidence or None)."""
    source, score = "text", None
    if _usable(pdf_text(data, name)):
        texts = [doc_text(data, name)]
        if os.getenv("EXTRACT_BACKEND", "").strip().lower() == "docling":
            texts.append(pdf_text(data, name))
    else:
        text, score = scan_text(data, name)
        texts, source = [text], "ocr"
    best = None
    for t in (t for t in texts if t.strip()):
        d = normalise(extract_heuristic(t))
        score_ = (len(validate(d)), -d.confidence)
        if best is None or score_ < best[0]:
            best = (score_, d)
    if best is None:
        raise ScannedImageError(SCANNED_MSG)
    return best[1], source, score


def process_file(name: str, data: bytes, force_heuristic: bool = False) -> dict:
    """Return a result dict for one file: {file, mode, model, source, document, needs_review, reasons}.

    mode is 'groq', 'gemini' or 'heuristic'; model is the model id that read the file ('regex' for heuristic);
    source is where the text came from: 'text' (PDF text layer), 'ocr' (scan or photo read by OCR; confidence is
    capped, see ocr_confidence) or 'vision' (Gemini read the page image); '' for a row that failed to extract."""
    prov = "none" if force_heuristic else active_provider()
    mode = prov if prov in ("groq", "gemini") else "heuristic"
    model, err, source, score = ("regex" if mode == "heuristic" else ""), None, "", None
    try:
        is_pdf = Path(name).suffix.lower() == ".pdf"
        if mode == "gemini":   # images go to the vision model as-is; PDFs also pass their text layer as a hint
            doc, model = extract_gemini(render_pages(data, name), doc_text(data, name) if is_pdf else ""), GEMINI_MODEL
            source = "vision"
        elif mode == "groq":
            doc, model, source, score = extract_groq(data, name, doc_text(data, name) if is_pdf else "")
        else:
            doc, source, score = heuristic_from_file(name, data)
        doc = normalise(doc)
        if source == "ocr":
            doc.confidence = ocr_confidence(doc.confidence, score)
    except Exception as e:
        doc, source = Document(confidence=0.0), ""
        err = str(e) if isinstance(e, ScannedImageError) else f"{type(e).__name__}: {' '.join(str(e).split())[:300]}"
    reasons = validate(doc) if not err else []
    if err:
        reasons.insert(0, f"extraction error: {err}")
    return {"file": name, "mode": mode, "model": model, "source": source, "document": doc,
            "needs_review": bool(reasons), "reasons": "; ".join(reasons)}


# ---------------------------------------------------------------- output
COLUMNS = ["file", "doc_type", "vendor", "document_number", "date", "currency", "subtotal", "tax", "tip", "total",
           "payment_method", "line_items", "confidence", "needs_review", "review_reason", "mode", "source"]
ITEM_COLUMNS = ["file", "vendor", "description", "qty", "unit_price", "amount"]
MONEY_COLS = ("subtotal", "tax", "tip", "total", "unit_price", "amount")


def to_row(r: dict) -> dict:
    d: Document = r["document"]
    items = " | ".join(f"{li.description} x{li.qty:g} @ {li.unit_price} = {li.amount}"
                       if li.qty is not None else f"{li.description} = {li.amount}" for li in d.line_items)
    return {"file": r["file"], "doc_type": d.doc_type, "vendor": d.vendor, "document_number": d.document_number,
            "date": d.date, "currency": d.currency, "subtotal": d.subtotal, "tax": d.tax, "tip": d.tip,
            "total": d.total, "payment_method": d.payment_method, "line_items": items, "confidence": d.confidence,
            "needs_review": r["needs_review"], "review_reason": r["reasons"], "mode": r["mode"],
            "source": r.get("source", "")}


def to_df(results: list) -> pd.DataFrame:
    return pd.DataFrame([to_row(r) for r in results], columns=COLUMNS)


def items_df(results: list) -> pd.DataFrame:
    rows = [{"file": r["file"], "vendor": r["document"].vendor, "description": li.description, "qty": li.qty,
             "unit_price": li.unit_price, "amount": li.amount} for r in results for li in r["document"].line_items]
    return pd.DataFrame(rows, columns=ITEM_COLUMNS)


def summarize(df: pd.DataFrame) -> dict:
    """Headline numbers: files processed, flagged, clean, and the total amount per currency."""
    flagged = int(df.needs_review.sum()) if len(df) else 0
    by_cur = df.dropna(subset=["total"]).groupby(df["currency"].fillna("?"))["total"].sum().round(2).to_dict()
    return {"files": len(df), "flagged": flagged, "clean": len(df) - flagged, "totals": by_cur}


def xlsx_bytes(df: pd.DataFrame, items: Optional[pd.DataFrame] = None) -> bytes:
    """Workbook with sheets: documents (all rows), needs_review (flagged only), line_items (one row per item)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    items = items if items is not None else pd.DataFrame(columns=ITEM_COLUMNS)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="documents")
        df[df.needs_review].to_excel(w, index=False, sheet_name="needs_review")
        items.to_excel(w, index=False, sheet_name="line_items")
        head_fill, flag_fill = PatternFill("solid", fgColor="1F3A5F"), PatternFill("solid", fgColor="FDE2E1")
        for ws in w.book.worksheets:
            headers = [c.value for c in ws[1]]
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            for i, h in enumerate(headers, 1):
                width = max([len(str(h))] + [len(str(c.value)) for c in ws[get_column_letter(i)][1:]] or [8])
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
                if h in MONEY_COLS:
                    for c in ws[get_column_letter(i)][1:]:
                        c.number_format = "#,##0.00"
            if "needs_review" in headers and ws.title != "line_items":
                col = headers.index("needs_review") + 1
                for row in ws.iter_rows(min_row=2):
                    if row[col - 1].value is True:
                        for c in row:
                            c.fill = flag_fill
    return buf.getvalue()


def push_to_sheet(df: pd.DataFrame) -> str:
    """Append rows to Google Sheet if configured. Returns a status message."""
    sa, sid = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON"), os.getenv("SHEET_ID")
    if not (sa and sid):
        return "Google Sheets not configured - skipped."
    import gspread
    gc = gspread.service_account(filename=sa)
    ws = gc.open_by_key(sid).sheet1
    ws.append_rows(df.fillna("").astype(str).values.tolist(), value_input_option="USER_ENTERED")
    return f"Appended {len(df)} rows to Google Sheet."
